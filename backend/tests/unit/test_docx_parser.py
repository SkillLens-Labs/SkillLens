from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.app.core.exceptions import ApplicationError
from backend.app.infrastructure.parsers.docx_parser import DOCXParser
from backend.app.infrastructure.parsers.models import (
    DocumentBlockType,
    DocumentType,
)


def create_docx(
    paragraphs: list[tuple[str, str | None]] | None = None,
    bullets: list[str] | None = None,
    tables: list[list[list[str]]] | None = None,
) -> bytes:
    """Create an in-memory DOCX document for testing."""
    document = Document()

    if paragraphs:
        for text, style in paragraphs:
            document.add_paragraph(text, style=style)

    if bullets:
        for text in bullets:
            paragraph = document.add_paragraph(text)

            p_pr = paragraph._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")

            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")

            num_id = OxmlElement("w:numId")
            num_id.set(qn("w:val"), "1")

            num_pr.append(ilvl)
            num_pr.append(num_id)
            p_pr.append(num_pr)

    if tables:
        for table_data in tables:
            if not table_data:
                continue

            column_count = max(len(row) for row in table_data)

            table = document.add_table(
                rows=len(table_data),
                cols=column_count,
            )

            for row_index, row_data in enumerate(table_data):
                for column_index, value in enumerate(row_data):
                    table.cell(row_index, column_index).text = value

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def test_parse_basic_docx() -> None:
    parser = DOCXParser()

    content = create_docx(
        paragraphs=[
            ("John Doe", None),
            ("Software Engineer", None),
            ("Experienced Python developer.", None),
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-001",
    )

    assert result.document_id == "doc-001"
    assert result.document_type == DocumentType.DOCX
    assert len(result.blocks) == 3

    assert "John Doe" in result.text
    assert "Software Engineer" in result.text
    assert "Experienced Python developer." in result.text


def test_parse_multiple_paragraphs_in_order() -> None:
    parser = DOCXParser()

    content = create_docx(
        paragraphs=[
            ("First paragraph", None),
            ("Second paragraph", None),
            ("Third paragraph", None),
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-order",
    )

    texts = [block.text for block in result.blocks]

    assert texts == [
        "First paragraph",
        "Second paragraph",
        "Third paragraph",
    ]


def test_heading_is_classified_correctly() -> None:
    parser = DOCXParser()

    content = create_docx(
        paragraphs=[
            ("Professional Experience", "Heading 1"),
            ("Software Engineer", None),
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-heading",
    )

    heading_blocks = [
        block
        for block in result.blocks
        if block.block_type == DocumentBlockType.HEADING
    ]

    assert len(heading_blocks) == 1
    assert heading_blocks[0].text == "Professional Experience"
    assert heading_blocks[0].style_name == "Heading 1"


def test_bullet_paragraph_is_classified_correctly() -> None:
    parser = DOCXParser()

    content = create_docx(
        bullets=[
            "Python",
            "FastAPI",
            "SQL",
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-bullets",
    )

    bullet_blocks = [
        block
        for block in result.blocks
        if block.block_type == DocumentBlockType.BULLET
    ]

    assert len(bullet_blocks) == 3

    assert [block.text for block in bullet_blocks] == [
        "Python",
        "FastAPI",
        "SQL",
    ]


def test_table_cells_are_extracted() -> None:
    parser = DOCXParser()

    content = create_docx(
        tables=[
            [
                ["Skill", "Level"],
                ["Python", "Advanced"],
                ["SQL", "Intermediate"],
            ]
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-table",
    )

    table_blocks = [
        block
        for block in result.blocks
        if block.block_type == DocumentBlockType.TABLE_CELL
    ]

    assert len(table_blocks) == 6

    extracted = [block.text for block in table_blocks]

    assert extracted == [
        "Skill",
        "Level",
        "Python",
        "Advanced",
        "SQL",
        "Intermediate",
    ]


def test_table_cell_provenance_is_preserved() -> None:
    parser = DOCXParser()

    content = create_docx(
        tables=[
            [
                ["A", "B"],
                ["C", "D"],
            ]
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-table-provenance",
    )

    table_blocks = [
        block
        for block in result.blocks
        if block.block_type == DocumentBlockType.TABLE_CELL
    ]

    assert len(table_blocks) == 4

    first = table_blocks[0]

    assert first.source.table_index == 0
    assert first.source.row_index == 0
    assert first.source.column_index == 0
    assert first.source.block_index is not None

    second = table_blocks[1]

    assert second.source.table_index == 0
    assert second.source.row_index == 0
    assert second.source.column_index == 1


def test_paragraph_provenance_is_preserved() -> None:
    parser = DOCXParser()

    content = create_docx(
        paragraphs=[
            ("First paragraph", None),
            ("Second paragraph", None),
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-paragraph-provenance",
    )

    assert len(result.blocks) == 2

    first = result.blocks[0]
    second = result.blocks[1]

    assert first.source.paragraph_index == 0
    assert second.source.paragraph_index == 1

    assert first.source.block_index is not None
    assert second.source.block_index is not None

    assert first.source.block_index < second.source.block_index


def test_empty_and_whitespace_paragraphs_are_ignored() -> None:
    parser = DOCXParser()

    content = create_docx(
        paragraphs=[
            ("Valid paragraph", None),
            ("", None),
            ("     ", None),
            ("Another valid paragraph", None),
        ]
    )

    result = parser.parse(
        content=content,
        document_id="doc-whitespace",
    )

    assert len(result.blocks) == 2

    assert [block.text for block in result.blocks] == [
        "Valid paragraph",
        "Another valid paragraph",
    ]

    for block in result.blocks:
        assert block.text.strip() != ""


def test_empty_docx_returns_no_blocks() -> None:
    parser = DOCXParser()

    content = create_docx()

    result = parser.parse(
        content=content,
        document_id="doc-empty",
    )

    assert result.document_type == DocumentType.DOCX
    assert result.blocks == []
    assert result.text == ""


def test_corrupted_docx_raises_application_error() -> None:
    parser = DOCXParser()

    corrupted_content = b"this-is-not-a-valid-docx"

    with pytest.raises(ApplicationError) as exc_info:
        parser.parse(
            content=corrupted_content,
            document_id="doc-corrupted",
        )

    assert exc_info.value.code == "INVALID_DOCX_DOCUMENT"


def test_empty_content_raises_application_error() -> None:
    parser = DOCXParser()

    with pytest.raises(ApplicationError) as exc_info:
        parser.parse(
            content=b"",
            document_id="doc-empty-content",
        )

    assert exc_info.value.code == "EMPTY_DOCUMENT"


def test_metadata_contains_document_structure_counts() -> None:
    parser = DOCXParser()

    content = create_docx(
        paragraphs=[
            ("Summary", "Heading 1"),
            ("Candidate summary", None),
            ("Experience", "Heading 1"),
        ],
        tables=[
            [
                ["Company", "Role"],
                ["Example Corp", "Developer"],
            ]
        ],
    )

    result = parser.parse(
        content=content,
        document_id="doc-metadata",
    )

    assert result.metadata["paragraph_count"] == 3
    assert result.metadata["table_count"] == 1
    assert result.metadata["parser"] == "python-docx"