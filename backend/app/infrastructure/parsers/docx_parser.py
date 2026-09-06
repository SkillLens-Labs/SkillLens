from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from backend.app.core.exceptions import ApplicationError
from backend.app.infrastructure.parsers.base import DocumentParser
from backend.app.infrastructure.parsers.models import (
    DocumentBlock,
    DocumentBlockType,
    DocumentType,
    ParsedDocument,
    SourceLocation,
)


class DOCXParser(DocumentParser):
    """Parser for extracting structured content from DOCX documents."""

    document_type = DocumentType.DOCX

    def parse(self, content: bytes, document_id: str) -> ParsedDocument:
        """
        Parse DOCX bytes into a ParsedDocument.

        The parser preserves:
        - paragraph ordering
        - heading information
        - bullet/list information
        - table-cell content
        - paragraph/table provenance
        """
        if not content:
            raise ApplicationError(
                code="EMPTY_DOCUMENT",
                message="The DOCX document is empty.",
            )

        try:
            document = Document(BytesIO(content))
        except Exception as exc:
            raise ApplicationError(
                code="INVALID_DOCX_DOCUMENT",
                message="The uploaded file is not a valid DOCX document.",
                details=str(exc),
            ) from exc

        try:
            blocks: list[DocumentBlock] = []

            block_index = 0
            paragraph_index = 0
            table_index = 0

            for element in document.element.body.iterchildren():
                if isinstance(element, CT_P):
                    paragraph = Paragraph(element, document)

                    block = self._parse_paragraph(
                        paragraph=paragraph,
                        paragraph_index=paragraph_index,
                        block_index=block_index,
                    )

                    paragraph_index += 1

                    if block is not None:
                        blocks.append(block)
                        block_index += 1

                elif isinstance(element, CT_Tbl):
                    table = Table(element, document)

                    table_blocks = self._parse_table(
                        table=table,
                        table_index=table_index,
                        starting_block_index=block_index,
                    )

                    table_index += 1

                    blocks.extend(table_blocks)
                    block_index += len(table_blocks)

            return ParsedDocument(
                document_id=document_id,
                document_type=DocumentType.DOCX,
                blocks=blocks,
                metadata={
                    "paragraph_count": paragraph_index,
                    "table_count": table_index,
                    "parser": "python-docx",
                },
            )

        except ApplicationError:
            raise

        except Exception as exc:
            raise ApplicationError(
                code="DOCX_PARSE_FAILED",
                message="Failed to parse the DOCX document.",
                details=str(exc),
            ) from exc

    @classmethod
    def _parse_paragraph(
        cls,
        paragraph: Paragraph,
        paragraph_index: int,
        block_index: int,
    ) -> DocumentBlock | None:
        """Convert a DOCX paragraph into a DocumentBlock."""
        text = paragraph.text.strip()

        if not text:
            return None

        style_name = paragraph.style.name if paragraph.style is not None else None

        block_type = cls._classify_paragraph(
            paragraph=paragraph,
            style_name=style_name,
        )

        source = SourceLocation(
            paragraph_index=paragraph_index,
            block_index=block_index,
        )

        return DocumentBlock(
            text=text,
            block_type=block_type,
            source=source,
            style_name=style_name,
            section=None,
            metadata={
                "is_list_item": cls._is_list_item(paragraph),
            },
        )

    @classmethod
    def _parse_table(
        cls,
        table: Table,
        table_index: int,
        starting_block_index: int,
    ) -> list[DocumentBlock]:
        """Extract non-empty table cells with row/column provenance."""
        blocks: list[DocumentBlock] = []
        block_index = starting_block_index

        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                text = cls._extract_cell_text(cell)

                if not text:
                    continue

                source = SourceLocation(
                    table_index=table_index,
                    row_index=row_index,
                    column_index=column_index,
                    block_index=block_index,
                )

                blocks.append(
                    DocumentBlock(
                        text=text,
                        block_type=DocumentBlockType.TABLE_CELL,
                        source=source,
                        style_name=None,
                        section=None,
                        metadata={
                            "table_index": table_index,
                            "row_index": row_index,
                            "column_index": column_index,
                        },
                    )
                )

                block_index += 1

        return blocks

    @staticmethod
    def _extract_cell_text(cell: _Cell) -> str:
        """
        Extract text from a table cell.

        Multiple paragraphs inside one cell are preserved as separate lines.
        """
        paragraphs = [
            paragraph.text.strip()
            for paragraph in cell.paragraphs
            if paragraph.text.strip()
        ]

        return "\n".join(paragraphs).strip()

    @staticmethod
    def _classify_paragraph(
        paragraph: Paragraph,
        style_name: str | None,
    ) -> DocumentBlockType:
        """Classify a DOCX paragraph using its native style/list metadata."""
        normalized_style = (style_name or "").strip().lower()

        if normalized_style.startswith("heading"):
            return DocumentBlockType.HEADING

        if DOCXParser._is_list_item(paragraph):
            return DocumentBlockType.BULLET

        return DocumentBlockType.PARAGRAPH

    @staticmethod
    def _is_list_item(paragraph: Paragraph) -> bool:
        """
        Detect whether a paragraph belongs to a DOCX numbering definition.

        python-docx does not expose a high-level `is_list_item` API, so the
        numbering XML is inspected directly.
        """
        p = paragraph._p
        p_pr = p.pPr

        if p_pr is None:
            return False

        num_pr = p_pr.numPr

        return num_pr is not None