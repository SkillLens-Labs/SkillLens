import io
import json

from docx import Document
from fastapi import APIRouter, File, Form, UploadFile, status

from backend.app.core.exceptions import ApplicationError
from backend.app.orchestration.analysis_orchestrator import AnalysisOrchestrator
from backend.app.orchestration.concrete_analysis_orchestrator import (
    ConcreteAnalysisOrchestrator,
)
from backend.app.schemas.requests import (
    JobDescriptionDocumentInput,
    ResumeAnalysisRequest,
    ResumeDocumentInput,
    ResumeJDAnalysisRequest,
)
from backend.app.schemas.responses import (
    AnalysisResponse,
    DeleteAnalysisResponse,
)

router = APIRouter(prefix="/analyses", tags=["analyses"])

_orchestrator: AnalysisOrchestrator = ConcreteAnalysisOrchestrator()


def _parse_request_payload(
    options: str | None,
    client_metadata: str | None,
) -> ResumeAnalysisRequest:
    parsed_options = None
    parsed_metadata = None

    if options is not None:
        try:
            parsed_options = json.loads(options)
        except json.JSONDecodeError as exc:
            raise ApplicationError(
                code="INVALID_ANALYSIS_OPTIONS",
                message="The options field must contain valid JSON.",
                field="options",
            ) from exc

    if client_metadata is not None:
        try:
            parsed_metadata = json.loads(client_metadata)
        except json.JSONDecodeError as exc:
            raise ApplicationError(
                code="INVALID_CLIENT_METADATA",
                message="The client_metadata field must contain valid JSON.",
                field="client_metadata",
            ) from exc

    try:
        payload = {}

        if parsed_options is not None:
            payload["options"] = parsed_options

        if parsed_metadata is not None:
            payload["client_metadata"] = parsed_metadata

        return ResumeAnalysisRequest.model_validate(payload)
    except ValueError as exc:
        raise ApplicationError(
            code="INVALID_ANALYSIS_REQUEST",
            message="The analysis request contains invalid fields.",
        ) from exc


def _text_as_docx(text: str, document_name: str) -> tuple[str, bytes, str]:
    if not text.strip():
        raise ApplicationError(
            code="EMPTY_DOCUMENT",
            message=f"The {document_name}_text field must not be empty.",
            field=f"{document_name}_text",
        )

    document = Document()

    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)

    output = io.BytesIO()
    document.save(output)

    return (
        f"{document_name}.docx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def _job_description_text_as_docx(
    text: str,
    document_name: str,
) -> tuple[str, bytes, str]:
    if not text.strip():
        raise ApplicationError(
            code="EMPTY_DOCUMENT",
            message=f"The {document_name}_text field must not be empty.",
            field=f"{document_name}_text",
        )

    document = Document()

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ApplicationError(
            code="EMPTY_DOCUMENT",
            message=f"The {document_name}_text field must not be empty.",
            field=f"{document_name}_text",
        )

    # Preserve the first line as the document title.
    document.add_heading(lines[0], level=1)

    # Put remaining plain text into a section the JD pipeline understands.
    if len(lines) > 1:
        document.add_heading("Responsibilities", level=2)
        for line in lines[1:]:
            document.add_paragraph(line)
    else:
        document.add_heading("Responsibilities", level=2)
        document.add_paragraph(lines[0])

    output = io.BytesIO()
    document.save(output)

    return (
        f"{document_name}.docx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


async def _read_document_input(
    file_value: UploadFile | None,
    text_value: str | None,
    *,
    document_name: str,
    is_job_description: bool = False,
) -> tuple[str, bytes, str]:
    has_file = file_value is not None
    has_text = text_value is not None

    if has_file and has_text:
        raise ApplicationError(
            code="INVALID_DOCUMENT_INPUT",
            message=(
                f"Provide exactly one of {document_name} file or "
                f"{document_name}_text, not both."
            ),
            field=document_name,
        )

    if not has_file and not has_text:
        raise ApplicationError(
            code="INVALID_DOCUMENT_INPUT",
            message=(
                f"Provide either {document_name} file or "
                f"{document_name}_text."
            ),
            field=document_name,
        )

    if text_value is not None:
        if is_job_description:
            return _job_description_text_as_docx(text_value, document_name)
        return _text_as_docx(text_value, document_name)

    assert file_value is not None

    content = await file_value.read()

    return (
        file_value.filename or f"{document_name}.docx",
        content,
        file_value.content_type
        or "application/octet-stream",
    )


@router.post(
    "/resume",
    response_model=AnalysisResponse,
    status_code=status.HTTP_200_OK,
)
async def analyze_resume(
    resume: UploadFile | None = File(default=None),
    resume_text: str | None = Form(default=None),
    options: str | None = Form(default=None),
    client_metadata: str | None = Form(default=None),
) -> AnalysisResponse:
    """Run synchronous resume-only analysis."""
    filename, content, content_type = await _read_document_input(
        resume,
        resume_text,
        document_name="resume",
    )

    request = _parse_request_payload(
        options=options,
        client_metadata=client_metadata,
    )

    result = _orchestrator.analyze_resume(
        ResumeDocumentInput(
            filename=filename,
            content=content,
            content_type=content_type,
        ),
        request,
    )

    return AnalysisResponse(data=result)


@router.post(
    "/resume-jd",
    response_model=AnalysisResponse,
    status_code=status.HTTP_200_OK,
)
async def analyze_resume_jd(
    resume: UploadFile | None = File(default=None),
    resume_text: str | None = Form(default=None),
    job_description: UploadFile | None = File(default=None),
    job_description_text: str | None = Form(default=None),
    options: str | None = Form(default=None),
    client_metadata: str | None = Form(default=None),
) -> AnalysisResponse:
    """Run synchronous resume + job-description analysis."""
    resume_filename, resume_content, resume_content_type = (
        await _read_document_input(
            resume,
            resume_text,
            document_name="resume",
        )
    )

    job_description_filename, job_description_content, job_description_content_type = (
        await _read_document_input(
            job_description,
            job_description_text,
            document_name="job_description",
            is_job_description=True,
        )
    )

    parsed_request = _parse_request_payload(
        options=options,
        client_metadata=client_metadata,
    )

    result = _orchestrator.analyze_resume_jd(
        ResumeDocumentInput(
            filename=resume_filename,
            content=resume_content,
            content_type=resume_content_type,
        ),
        JobDescriptionDocumentInput(
            filename=job_description_filename,
            content=job_description_content,
            content_type=job_description_content_type,
        ),
        ResumeJDAnalysisRequest(
            options=parsed_request.options,
            client_metadata=parsed_request.client_metadata,
        ),
    )

    return AnalysisResponse(data=result)


@router.get(
    "/{analysis_id}",
    response_model=AnalysisResponse,
    status_code=status.HTTP_200_OK,
)
async def get_analysis(analysis_id: str) -> AnalysisResponse:
    """Retrieve a completed analysis result."""
    result = _orchestrator.get_analysis(analysis_id)

    if result is None:
        raise ApplicationError(
            code="ANALYSIS_NOT_FOUND",
            message=f"Analysis '{analysis_id}' was not found.",
            field="analysis_id",
        )

    return AnalysisResponse(data=result)


@router.delete(
    "/{analysis_id}",
    response_model=DeleteAnalysisResponse,
    status_code=status.HTTP_200_OK,
)
async def delete_analysis(analysis_id: str) -> DeleteAnalysisResponse:
    """Delete an existing analysis result."""
    deleted = _orchestrator.delete_analysis(analysis_id)

    if not deleted:
        raise ApplicationError(
            code="ANALYSIS_NOT_FOUND",
            message=f"Analysis '{analysis_id}' was not found.",
            field="analysis_id",
        )

    return DeleteAnalysisResponse(
        analysis_id=analysis_id,
        deleted=True,
        message="Analysis deleted successfully.",
    )